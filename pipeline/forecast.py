"""Forecast for history-matched parameter sets, and writing it into the organiser's template.

Method A+ = Approach A (proxy history → Arps/exponential decline) with the forecast tied to the observed field:
- each phase's forecast rate is scaled by (last observed rate / last proxy rate), so the matched model supplies
  the decline *shape* and the field supplies the starting *level*;
- cumulatives start from the observed 2008-01-01 cumulatives.
This removes the jump at the history/forecast boundary caused by the late-time water-cut mismatch. On field
backtests (pipeline/main.py) it beats the unanchored forecast for oil and on average for water.
"""
import numpy as np
import openpyxl
import pandas as pd

from .config import DECLINE_WINDOW, FC_DATES
from .proxy import decline_forecast, integrate_cum

TEMPLATE_COLS = {'gas': 2, 'oil': 3, 'water': 4}  # 1-based Excel columns in each template sheet


def forecast(proxy, param_sets, hist_dates, obs_rate, obs_cum, corr, fc_dates=FC_DATES, anchor=True):
    """param_sets: (n, 4) raw parameters. obs_rate/obs_cum: observed history (simulator convention) at hist_dates.
    Returns dict with rate/cum forecasts (n, F, 3) and the proxy history (n, T, 3)."""
    r_hist, c_hist = proxy.predict(param_sets)
    rate, cum, fits = decline_forecast(r_hist, c_hist, hist_dates, fc_dates, DECLINE_WINDOW, corr)
    if anchor:
        rate = rate * (obs_rate[-1] / np.maximum(r_hist[:, -1], 1e-9))[:, None, :]
        cum = integrate_cum(rate, np.repeat(obs_cum[-1][None], len(rate), 0), hist_dates[-1], fc_dates, corr)
    return {'rate': rate, 'cum': cum, 'rate_hist': r_hist, 'cum_hist': c_hist, 'fits': fits}


def write_template(cum, template_path, out_path):
    """Fill the template's value cells in place (dates and formatting untouched). cum: (10, 80, 3) [oil, gas, water]."""
    wb = openpyxl.load_workbook(template_path)
    assert len(wb.sheetnames) == len(cum), f'template has {len(wb.sheetnames)} sheets, got {len(cum)} cases'
    for k, ws in enumerate(wb.worksheets):
        dates = pd.to_datetime([ws.cell(r, 1).value for r in range(2, ws.max_row + 1)], format='mixed')
        assert pd.DatetimeIndex(dates).equals(FC_DATES), f'{ws.title}: template dates differ from the forecast dates'
        for t in range(len(FC_DATES)):
            for v, name in enumerate(['oil', 'gas', 'water']):
                ws.cell(t + 2, TEMPLATE_COLS[name], float(cum[k, t, v]))
    wb.save(out_path)
